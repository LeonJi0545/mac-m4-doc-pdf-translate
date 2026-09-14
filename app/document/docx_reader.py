"""DOCX 解析（方案 §5.1）。

结构化解析，**不**把文档转成纯文本 —— 翻译完要写回原结构（见 docx_writer）。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

from app.core.errors import ParseError
from app.document.docx_walk import (
    iter_paragraph_entries,
    run_has_image,
    runs_text,
    split_paragraph,
    table_coords,
)
from app.document.model import Block, BlockType, ParsedDocument, TableRef

_HEADING_LEVEL = re.compile(r"heading\s*(\d+)", re.I)


def _classify(paragraph: Any, key: str) -> tuple[BlockType, int | None]:
    style_name = ""
    try:
        style_name = (paragraph.style.name or "") if paragraph.style else ""
    except (AttributeError, KeyError):
        style_name = ""

    if table_coords(key):
        return BlockType.TABLE_CELL, None

    lowered = style_name.lower()
    if lowered.startswith("title"):
        return BlockType.HEADING, 1
    m = _HEADING_LEVEL.search(style_name)
    if m:
        return BlockType.HEADING, int(m.group(1))
    if "caption" in lowered:
        return BlockType.CAPTION, None
    if "list" in lowered:
        return BlockType.LIST_ITEM, None
    return BlockType.PARAGRAPH, None


def parse_docx(path: Path) -> ParsedDocument:
    """把 .docx 解析成统一的 Block 序列。"""
    try:
        document = Document(str(path))
    except Exception as exc:  # python-docx 对损坏文件抛的类型很杂
        raise ParseError(f"DOCX 解析失败 ({path}): {exc}") from exc

    blocks: list[Block] = []
    image_count = 0
    seen_keys: set[str] = set()

    for key, paragraph in iter_paragraph_entries(document):
        # 合并单元格会在多个坐标上重复出现同一个段落对象，去重避免翻译两遍
        if key in seen_keys:
            continue
        seen_keys.add(key)

        parts = split_paragraph(paragraph)
        block_type, level = _classify(paragraph, key)

        coords = table_coords(key)
        table_ref = (
            TableRef(table_id=coords[0], row=coords[1], col=coords[2]) if coords else None
        )

        image_count += sum(1 for r in parts.text_runs if run_has_image(r))
        for group in parts.hyperlink_runs:
            image_count += sum(1 for r in group if run_has_image(r))

        # 正文（不含超链接文字）
        main_text = runs_text([r for r in parts.text_runs if not run_has_image(r)])
        if main_text.strip():
            blocks.append(
                Block(
                    id=f"{key}#main",
                    type=block_type,
                    text=main_text,
                    level=level,
                    table=table_ref,
                    anchor={"kind": "docx", "key": key, "part": "main"},
                )
            )

        # 每个超链接的显示文字单独成块 —— 这样译文不会把整段变成链接，
        # 也不会留下空链接，URL 全程没被碰过（它在 relationship 里）。
        for idx, group in enumerate(parts.hyperlink_runs):
            link_text = runs_text([r for r in group if not run_has_image(r)])
            if link_text.strip():
                blocks.append(
                    Block(
                        id=f"{key}#link{idx}",
                        type=block_type,
                        text=link_text,
                        level=level,
                        table=table_ref,
                        anchor={
                            "kind": "docx",
                            "key": key,
                            "part": "hyperlink",
                            "index": idx,
                        },
                    )
                )

    return ParsedDocument(
        source_path=path,
        kind="docx",
        blocks=blocks,
        meta={"image_count": image_count},
    )


def count_structure(path: Path) -> dict[str, int]:
    """统计段落 / 表格 / 标题 / 图片数量，供 QA 做源与输出的比对（方案 §14）。"""
    document = Document(str(path))
    paragraphs = 0
    headings = 0
    images = 0
    seen: set[str] = set()
    tables: set[str] = set()

    for key, paragraph in iter_paragraph_entries(document):
        if key in seen:
            continue
        seen.add(key)
        paragraphs += 1
        block_type, _ = _classify(paragraph, key)
        if block_type is BlockType.HEADING:
            headings += 1
        coords = table_coords(key)
        if coords:
            tables.add(coords[0])
        parts = split_paragraph(paragraph)
        images += sum(1 for r in parts.text_runs if run_has_image(r))
        for group in parts.hyperlink_runs:
            images += sum(1 for r in group if run_has_image(r))

    return {
        "paragraphs": paragraphs,
        "tables": len(tables),
        "headings": headings,
        "images": images,
    }
