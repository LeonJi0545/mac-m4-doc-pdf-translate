"""现场构造测试用 DOCX —— 不往仓库里放二进制 fixture。"""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

# 1x1 透明 PNG，内联成常量避免引入外部文件
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """插入一个真实的外部超链接（URL 存在 relationship 里）。"""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_rich_docx(path: Path) -> Path:
    """构造一份覆盖 §5.1 全部保留项的文档。

    标题 / 段落 / 列表 / 表格 / 嵌套表格 / 页眉页脚 / 图片 / 超链接 / 多 run 混合格式。
    """
    doc = Document()

    doc.add_heading("Quarterly Operations Report", level=1)
    doc.add_paragraph("The invoice totals EUR 1,250.00 and is due on 2024-03-15.")

    # 多 run 的段落（模拟句中加粗）
    p = doc.add_paragraph("Contact ")
    p.add_run("support").bold = True
    p.add_run(" before 2024-04-01.")

    doc.add_heading("Details", level=2)
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")

    # 带超链接的段落
    link_para = doc.add_paragraph("See ")
    add_hyperlink(link_para, "https://example.com/policy", "the policy page")
    link_para.add_run(" for details.")

    # 图片
    doc.add_picture(io.BytesIO(PNG_1X1), width=Inches(0.5))

    # 表格（含嵌套表格）
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "Lombardia"
    inner = table.cell(1, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "EUR 980.00"

    # 页眉页脚
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Internal use only"
    section.footer.paragraphs[0].text = "Page footer note"

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def build_simple_docx(path: Path, paragraphs: list[str]) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def count_drawings(path: Path) -> int:
    """统计文档里的图片元素（w:drawing + w:pict）。"""
    doc = Document(str(path))
    body = doc.element.body
    total = len(body.findall(f".//{qn('w:drawing')}")) + len(body.findall(f".//{qn('w:pict')}"))
    for section in doc.sections:
        for part in (section.header, section.footer):
            el = part._element
            total += len(el.findall(f".//{qn('w:drawing')}")) + len(
                el.findall(f".//{qn('w:pict')}")
            )
    return total


def hyperlink_targets(path: Path) -> list[str]:
    """取出文档里全部外部超链接的 URL。"""
    doc = Document(str(path))
    rels = doc.part.rels
    return sorted(
        rel.target_ref
        for rel in rels.values()
        if "hyperlink" in rel.reltype and rel.is_external
    )
