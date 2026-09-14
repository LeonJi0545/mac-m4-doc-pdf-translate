"""M2 测试：DOCX 解析 → 翻译 → 写回原结构（AC-2.1 ~ AC-2.4）。"""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

from app.document.docx_reader import count_structure, parse_docx
from app.document.docx_writer import apply_translations, write_docx
from app.document.model import BlockType
from tests.fixtures.docx_factory import (
    build_rich_docx,
    count_drawings,
    hyperlink_targets,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _translate_all(parsed, prefix: str = "[zh]") -> None:
    apply_translations(
        parsed, {b.id: f"{prefix}{b.text}" for b in parsed.blocks if b.translatable}
    )


def test_parses_all_structure(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    parsed = parse_docx(src)

    texts = [b.text for b in parsed.blocks]
    assert any("Quarterly Operations Report" in t for t in texts)
    assert any("EUR 1,250.00" in t for t in texts)
    assert any("Internal use only" in t for t in texts)  # 页眉
    assert any("Page footer note" in t for t in texts)  # 页脚
    assert any("Lombardia" in t for t in texts)  # 表格
    assert any("EUR 980.00" in t for t in texts)  # 嵌套表格

    assert any(b.type is BlockType.HEADING for b in parsed.blocks)
    assert any(b.type is BlockType.TABLE_CELL for b in parsed.blocks)
    assert any(b.type is BlockType.LIST_ITEM for b in parsed.blocks)


def test_hyperlink_text_is_its_own_block(tmp_path: Path) -> None:
    """超链接文字单独成块 —— 否则要么整段变成链接、要么留下空链接。"""
    src = build_rich_docx(tmp_path / "rich.docx")
    parsed = parse_docx(src)
    link_blocks = [b for b in parsed.blocks if b.anchor.get("part") == "hyperlink"]
    assert [b.text for b in link_blocks] == ["the policy page"]

    # 同一段落的正文块不应重复包含链接文字
    main = [b for b in parsed.blocks if b.id.startswith(link_blocks[0].anchor["key"] + "#main")]
    assert main and "the policy page" not in main[0].text


# ── AC-2.1 结构计数一致 ───────────────────────────────────────────────────


def test_structure_counts_preserved(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)

    before = count_structure(src)
    after = count_structure(out)
    assert before == after, f"结构计数变了: {before} -> {after}"
    assert before["tables"] >= 2  # 主表 + 嵌套表
    assert before["headings"] >= 2


# ── AC-2.2 图片保真 ───────────────────────────────────────────────────────


def test_images_survive_writeback(tmp_path: Path) -> None:
    """写回时清空 run 是常见做法，但含图片的 run 一碰图片就没了。"""
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)

    assert count_drawings(src) >= 1
    assert count_drawings(out) == count_drawings(src)


# ── AC-2.3 超链接 URL 不变 ────────────────────────────────────────────────


def test_hyperlink_url_untouched(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)

    assert hyperlink_targets(src) == ["https://example.com/policy"]
    assert hyperlink_targets(out) == hyperlink_targets(src)


def test_hyperlink_display_text_is_translated(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)

    all_text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "[zh]the policy page" in all_text
    # 链接文字没有被复制到正文里
    assert all_text.count("the policy page") == 1


# ── AC-2.4 原文件永不修改 ─────────────────────────────────────────────────


def test_source_file_never_modified(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    before = _sha256(src)
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, tmp_path / "out.docx")
    assert _sha256(src) == before


def test_translation_actually_lands(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)

    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "[zh]Quarterly Operations Report" in body
    header = doc.sections[0].header.paragraphs[0].text
    assert header.startswith("[zh]")
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert any(c.startswith("[zh]Lombardia") for c in cells)


def test_multi_run_paragraph_is_one_block(tmp_path: Path) -> None:
    """句中加粗的段落应作为**一个**翻译单元，不该被 run 边界切碎（方案 §10）。"""
    src = build_rich_docx(tmp_path / "rich.docx")
    parsed = parse_docx(src)
    matches = [b for b in parsed.blocks if "support" in b.text]
    assert len(matches) == 1
    assert matches[0].text == "Contact support before 2024-04-01."


def test_no_partial_file_left_behind(tmp_path: Path) -> None:
    src = build_rich_docx(tmp_path / "rich.docx")
    out = tmp_path / "out.docx"
    parsed = parse_docx(src)
    _translate_all(parsed)
    write_docx(parsed, src, out)
    assert not list(tmp_path.glob("*.partial"))
